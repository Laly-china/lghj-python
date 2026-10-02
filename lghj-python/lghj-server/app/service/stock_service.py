"""股票基础数据业务（Excel 导入）。

对应复现原 Java：
- feng-lghj/lghj-server/src/main/java/com/lghj/service/IStockService.java
- feng-lghj/lghj-server/src/main/java/com/lghj/service/impl/StockServiceImpl.java
- feng-lghj/lghj-server/src/main/java/com/lghj/listener/StockExcelListener.java（EasyExcel 逐行转换）

Excel：A股详细数据.xlsx，列索引照抄 StockExcel（0代码/1名称/3简称/4总股本/
5流通股/6总市值/7流通市值/8行业/9上市日期，索引 2 跳过）。
EasyExcel 默认首行为表头（headRowNumber=1），解析从第 2 行开始；每 100 条一批入库。
"""

from __future__ import annotations

import logging
import os
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from sqlalchemy.orm import Session

from app.config import settings
from app.mapper import stock_basic_mapper
from app.pojo.entity import StockBasic

logger = logging.getLogger(__name__)

# 原 StockServiceImpl.DEFAULT_EXCEL_PATH = "excel/A股详细数据.xlsx"
DEFAULT_EXCEL_PATH = os.path.join("excel", "A股详细数据.xlsx")
BATCH_COUNT = 100  # 原 StockExcelListener.BATCH_COUNT


def convert_to_entity(row: tuple | list) -> StockBasic | None:
    """Excel 行转实体（对应 StockExcelListener.convertToEntity）。

    row 为按列索引取出的单元格值（None 表示空单元格）。
    symbol/name 为空的行直接跳过（对应 invoke 中的 hasText 判断）。
    """
    def cell(idx: int) -> str | None:
        if idx >= len(row):
            return None
        value = row[idx]
        if value is None:
            return None
        return str(value)

    symbol = normalize_code(cell(0))
    name = cell(1)
    if not symbol or not name:
        return None

    return StockBasic(
        symbol=symbol,
        name=name,
        short_name=cell(3),
        total_shares=_parse_long(cell(4)),
        float_shares=_parse_long(cell(5)),
        total_market_cap=_parse_long(cell(6)),
        float_market_cap=_parse_long(cell(7)),
        industry=cell(8),
        market_type=determine_market_type(symbol),
        list_date=parse_list_date(cell(9)),
        is_deleted=0,
    )


def normalize_code(code: str | None) -> str | None:
    """股票代码归一化（照抄 normalizeCode：取数字，超 6 位截前 6 位）。"""
    if not code:
        return code
    digits = re.sub(r"[^0-9]", "", code)
    if len(digits) > 6:
        digits = digits[:6]
    return digits if len(digits) == 6 else code.strip()


def determine_market_type(code: str | None) -> int:
    """市场类型判定（照抄 determineMarketType）。"""
    if code is None or not re.fullmatch(r"\d{6}", code):
        return 0
    if code.startswith(("600", "601", "603", "605")):
        return 1
    if code.startswith("688"):
        return 4
    if code.startswith(("000", "001", "002", "003")):
        return 2
    if code.startswith(("300", "301")):
        return 3
    if code.startswith(("83", "87", "88")):
        return 5
    if code.startswith("43"):
        return 6
    return 0


def parse_list_date(list_date_str: str | None) -> date | None:
    """上市日期解析（照抄 parseListDate：剔除非数字取前 8 位，yyyyMMdd）。"""
    if not list_date_str:
        return None
    cleaned = re.sub(r"[^0-9]", "", list_date_str)
    if len(cleaned) > 8:
        cleaned = cleaned[:8]
    if len(cleaned) != 8:
        return None
    try:
        return datetime.strptime(cleaned, "%Y%m%d").date()
    except ValueError:
        logger.warning("无法解析上市时间: %s", list_date_str)
        return None


def _parse_long(value: str | None) -> int | None:
    """股本/市值解析（照抄 parseLong：剔除逗号与空白，BigDecimal→long）。"""
    if not value:
        return None
    try:
        return int(Decimal(re.sub(r"[,\s]", "", value)))
    except (InvalidOperation, ValueError):
        logger.warning("无法解析为 Long: %s", value)
        return None


def resolve_excel_path() -> str:
    """Excel 路径解析（照抄 resolveExcelPath 优先级）。

    1. 配置/环境变量 LGHJ_STOCK_EXCEL_PATH（对应 -Dlghj.stock.excel-path 与环境变量）
    2. 相对路径 excel/A股详细数据.xlsx
    3. feng-lghj/excel/A股详细数据.xlsx
    """
    configured = settings.stock_excel_path or os.getenv("LGHJ_STOCK_EXCEL_PATH", "")
    if configured:
        return configured

    default_path = os.path.abspath(DEFAULT_EXCEL_PATH)
    if os.path.exists(default_path):
        return default_path

    module_path = os.path.abspath(os.path.join("feng-lghj", DEFAULT_EXCEL_PATH))
    if os.path.exists(module_path):
        return module_path

    return default_path


def read_excel_rows(excel_path: str) -> list[StockBasic]:
    """读取 Excel 并逐行转换（对应 EasyExcel.read(...).sheet().doRead + Listener）。

    openpyxl 读取，首行视为表头跳过（对齐 EasyExcel headRowNumber=1 默认行为）。
    """
    from openpyxl import load_workbook

    logger.info("Start reading stock Excel: %s", excel_path)
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    sheet = workbook.active

    stock_list: list[StockBasic] = []
    try:
        for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
            if row_index == 0:  # 跳过表头行
                continue
            entity = convert_to_entity(row)
            if entity is None:
                continue
            stock_list.append(entity)
    finally:
        workbook.close()
    return stock_list


def import_stock_basic(db: Session, excel_path: str | None = None) -> str:
    """从 Excel 导入 A 股基础信息（对应 StockServiceImpl.importStockBasic）。

    每 BATCH_COUNT=100 条一批入库；整体失败返回 "导入失败: ..."（原实现吞异常
    并以 Result.success 包裹该文案，此处保持一致）。
    """
    try:
        path = excel_path or resolve_excel_path()
        stock_list = read_excel_rows(path)

        batch: list[StockBasic] = []
        for stock in stock_list:
            batch.append(stock)
            if len(batch) >= BATCH_COUNT:
                stock_basic_mapper.insert_batch(db, batch)
                logger.info("开始保存 %d 条股票基础数据", len(batch))
                batch = []
        if batch:
            stock_basic_mapper.insert_batch(db, batch)
            logger.info("开始保存 %d 条股票基础数据", len(batch))

        db.commit()
        logger.info("股票基础数据 Excel 解析完成")
        return "成功导入股票数据"
    except Exception as exc:  # noqa: BLE001 对齐原 try/catch 行为
        db.rollback()
        logger.error("Failed to import stock Excel: %s", exc, exc_info=exc)
        return f"导入失败: {exc}"
